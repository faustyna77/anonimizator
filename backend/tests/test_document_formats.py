from io import BytesIO

import pymupdf
from docx import Document
import pytest

from backend.app.anonymization import CustomRegexRule, DocumentAnonymizer, PhraseRule
from backend.app.document_formats import DocumentProcessingError, anonymize_docx, anonymize_pdf
from backend.tests.pii_manifest import (
    CONTROLLED_NEAR_MISSES,
    EXPECTED_MAPPING,
    EXPECTED_MARKERS,
    SOURCE_PII_VALUES,
    make_pii_docx,
    make_pii_pdf,
    read_docx_text,
    read_pdf_text,
)


SYNTHETIC_IDENTIFIERS = (
    "44051401458 8567346215 synthetic.person@example.test 192598474 +48 123 456 789 "
    "11 1140 2004 0000 3002 1800 0002 PL61 1090 1014 0000 0712 1981 2874"
)
SOURCE_IDENTIFIERS = (
    "44051401458",
    "8567346215",
    "synthetic.person@example.test",
    "192598474",
    "+48 123 456 789",
    "11 1140 2004 0000 3002 1800 0002",
    "PL61 1090 1014 0000 0712 1981 2874",
)


def _synthetic_pdf(identifier_lines: list[str], *, include_static_text: bool = True) -> bytes:
    document = pymupdf.open()
    first_page = document.new_page(width=595, height=842)
    first_page.draw_rect(pymupdf.Rect(35, 35, 250, 65), color=(0, 0, 1), fill=(0.9, 0.9, 1))
    for index, identifier in enumerate(identifier_lines):
        first_page.insert_text((40, 105 + index * 24), identifier, fontsize=10)
    if include_static_text:
        first_page.insert_text((40, 280), "synthetic layout text remains", fontsize=10)

    second_page = document.new_page(width=595, height=842)
    if include_static_text:
        second_page.insert_text((40, 100), "second page survives unchanged", fontsize=10)

    output = BytesIO()
    document.save(output)
    document.close()
    return output.getvalue()


def _docx_text(document: Document) -> str:
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        + [paragraph.text for paragraph in document.sections[0].header.paragraphs]
        + [paragraph.text for paragraph in document.sections[0].footer.paragraphs]
    )


def test_docx_anonymization_uses_the_injected_office_rule_anonymizer():
    source = Document()
    source.add_paragraph("Jan Kowalski i Jan Kowalski")
    source_bytes = BytesIO()
    source.save(source_bytes)

    output, mapping = anonymize_docx(
        source_bytes.getvalue(),
        DocumentAnonymizer(phrase_rules=(PhraseRule(rule_id="office-rule", phrase="Jan Kowalski", marker_label="KLIENT"),)),
    )

    assert Document(BytesIO(output)).paragraphs[0].text == "[KLIENT_1] i [KLIENT_1]"
    assert mapping == {"[KLIENT_1]": "Jan Kowalski"}


def test_docx_anonymization_preserves_runs_and_handles_body_table_header_and_footer():
    source = Document()
    body = source.add_paragraph()
    retained_run = body.add_run("Retained ")
    retained_run.italic = True
    identifier_run = body.add_run("440514")
    identifier_run.bold = True
    continuation_run = body.add_run("01458")
    continuation_run.bold = True
    body.add_run(" 8567346215 synthetic.person@example.test +48 123 456 789 ")
    body.add_run("PL61 1090 1014 0000 0712 1981 2874")

    table_paragraph = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    table_paragraph.add_run(SYNTHETIC_IDENTIFIERS)
    section = source.sections[0]
    section.header.paragraphs[0].add_run(SYNTHETIC_IDENTIFIERS)
    section.footer.paragraphs[0].add_run(SYNTHETIC_IDENTIFIERS)
    source_bytes = BytesIO()
    source.save(source_bytes)

    output, mapping = anonymize_docx(source_bytes.getvalue(), DocumentAnonymizer())
    transformed = Document(BytesIO(output))
    text = _docx_text(transformed)
    transformed_body = transformed.paragraphs[0]

    assert output.startswith(b"PK")
    assert transformed_body.runs[0].text == "Retained "
    assert transformed_body.runs[0].italic is True
    assert transformed_body.runs[1].text == "[PESEL_1]"
    assert transformed_body.runs[1].bold is True
    assert "[PESEL_1]" in text
    assert "[NIP_1]" in text
    assert "[EMAIL_1]" in text
    assert "[REGON_1]" in text
    assert "[TELEFON_1]" in text
    assert "[NRB_1]" in text
    assert "[IBAN_1]" in text
    assert not any(identifier in text for identifier in SOURCE_IDENTIFIERS)
    assert set(mapping) == {"[PESEL_1]", "[NIP_1]", "[EMAIL_1]", "[REGON_1]", "[TELEFON_1]", "[NRB_1]", "[IBAN_1]"}


def test_pdf_anonymization_uses_value_only_regex_for_repeated_text():
    source = _synthetic_pdf(["SPRAWA-2026", "prefix SPRAWA-2026 suffix"])
    anonymizer = DocumentAnonymizer(
        regex_rules=(
            CustomRegexRule(rule_id="office-rule", pattern=r"SPRAWA-\d{4}", marker_label="NUMER_SPRAWY"),
        )
    )

    output, mapping = anonymize_pdf(source, anonymizer)
    document = pymupdf.open(stream=output, filetype="pdf")
    try:
        text = document[0].get_text("text")
    finally:
        document.close()

    assert text.count("[NUMER_SPRAWY_1]") == 2
    assert "prefix" in text and "suffix" in text
    assert "SPRAWA-2026" not in text
    assert mapping == {"[NUMER_SPRAWY_1]": "SPRAWA-2026"}


def test_pdf_anonymization_redacts_in_place_and_preserves_pages_layout_and_static_content():
    source = _synthetic_pdf(list(SOURCE_IDENTIFIERS))

    output, mapping = anonymize_pdf(source, DocumentAnonymizer())

    document = pymupdf.open(stream=output, filetype="pdf")
    text = "\n".join(page.get_text("text") for page in document)
    layout_rectangle = pymupdf.Rect(35, 35, 250, 65)
    try:
        assert len(document) == 2
        assert document[0].rect == pymupdf.Rect(0, 0, 595, 842)
        assert any(drawing["rect"].intersects(layout_rectangle) for drawing in document[0].get_drawings())
    finally:
        document.close()

    assert "[PESEL_1]" in text
    assert "[NIP_1]" in text
    assert "[EMAIL_1]" in text
    assert "[REGON_1]" in text
    assert "[TELEFON_1]" in text
    assert "[NRB_1]" in text
    assert "[IBAN_1]" in text
    assert "synthetic layout text remains" in text
    assert "second page survives unchanged" in text
    assert not any(identifier in text for identifier in SOURCE_IDENTIFIERS)
    assert not any(identifier.encode("utf-8") in output for identifier in SOURCE_IDENTIFIERS)
    assert set(mapping) == {"[PESEL_1]", "[NIP_1]", "[EMAIL_1]", "[REGON_1]", "[TELEFON_1]", "[NRB_1]", "[IBAN_1]"}


@pytest.mark.parametrize(
    ("source_factory", "anonymize", "text_reader"),
    (
        pytest.param(make_pii_docx, anonymize_docx, read_docx_text, id="docx"),
        pytest.param(make_pii_pdf, anonymize_pdf, read_pdf_text, id="pdf"),
    ),
)
def test_real_format_adapters_match_the_independent_pii_manifest(
    source_factory,
    anonymize,
    text_reader,
):
    output, mapping = anonymize(source_factory(), DocumentAnonymizer())
    text = text_reader(output)

    assert all(marker in text for marker in EXPECTED_MARKERS)
    assert not any(value in text for value in SOURCE_PII_VALUES)
    assert all(near_miss in text for near_miss in CONTROLLED_NEAR_MISSES)
    assert mapping == EXPECTED_MAPPING


def test_pdf_without_text_layer_returns_a_controlled_ocr_error():
    source = _synthetic_pdf([], include_static_text=False)

    try:
        anonymize_pdf(source, DocumentAnonymizer())
    except DocumentProcessingError as error:
        assert str(error) == "PDF text layer is unavailable; OCR is not supported in this MVP."
    else:
        raise AssertionError("a PDF without text unexpectedly processed")
