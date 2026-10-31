"""Independent synthetic PII oracle and real-document factories for backend tests."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document
import pymupdf


@dataclass(frozen=True)
class PIIManifestEntry:
    name: str
    marker: str
    first_spelling: str
    canonical_variant: str
    near_miss: str


REGON_NEAR_MISS = "192598475"


PII_MANIFEST = (
    PIIManifestEntry(
        name="PESEL",
        marker="[PESEL_1]",
        first_spelling="44051401458",
        canonical_variant="44051401458",
        near_miss="44051401459",
    ),
    PIIManifestEntry(
        name="NIP",
        marker="[NIP_1]",
        first_spelling="856-73-46-215",
        canonical_variant="8567346215",
        near_miss="856-73-46-214",
    ),
    PIIManifestEntry(
        name="EMAIL",
        marker="[EMAIL_1]",
        first_spelling="Synthetic.Person@Example.Test",
        canonical_variant="synthetic.person@example.test",
        near_miss="synthetic.person@example",
    ),
    PIIManifestEntry(
        name="REGON",
        marker="[REGON_1]",
        first_spelling="192598474",
        canonical_variant="192598474",
        near_miss=REGON_NEAR_MISS,
    ),
    PIIManifestEntry(
        name="TELEFON",
        marker="[TELEFON_1]",
        first_spelling="+48 123 456 789",
        canonical_variant="+48-123-456-789",
        near_miss="+48 123 456 78",
    ),
    PIIManifestEntry(
        name="NRB",
        marker="[NRB_1]",
        first_spelling="11 1140 2004 0000 3002 1800 0002",
        canonical_variant="11114020040000300218000002",
        near_miss="11 1140 2004 0000 3002 1800 0003",
    ),
    PIIManifestEntry(
        name="IBAN",
        marker="[IBAN_1]",
        first_spelling="PL61 1090 1014 0000 0712 1981 2874",
        canonical_variant="pl61109010140000071219812874",
        near_miss="PL61 1090 1014 0000 0712 1981 2875",
    ),
)

EXPECTED_MARKERS = (
    "[PESEL_1]",
    "[NIP_1]",
    "[EMAIL_1]",
    "[REGON_1]",
    "[TELEFON_1]",
    "[NRB_1]",
    "[IBAN_1]",
)
EXPECTED_MAPPING = {
    "[PESEL_1]": "44051401458",
    "[NIP_1]": "856-73-46-215",
    "[EMAIL_1]": "Synthetic.Person@Example.Test",
    "[REGON_1]": "192598474",
    "[TELEFON_1]": "+48 123 456 789",
    "[NRB_1]": "11 1140 2004 0000 3002 1800 0002",
    "[IBAN_1]": "PL61 1090 1014 0000 0712 1981 2874",
}
SOURCE_PII_VALUES = (
    "44051401458",
    "856-73-46-215",
    "8567346215",
    "Synthetic.Person@Example.Test",
    "synthetic.person@example.test",
    "192598474",
    "+48 123 456 789",
    "+48-123-456-789",
    "11 1140 2004 0000 3002 1800 0002",
    "11114020040000300218000002",
    "PL61 1090 1014 0000 0712 1981 2874",
    "pl61109010140000071219812874",
)
CONTROLLED_NEAR_MISSES = (
    "44051401459",
    "856-73-46-214",
    "synthetic.person@example",
    "+48 123 456 78",
    "11 1140 2004 0000 3002 1800 0003",
    "PL61 1090 1014 0000 0712 1981 2875",
)

PII_DOCUMENT_LINES = (
    "PESEL: 44051401458; PESEL canonical repeat: 44051401458.",
    "NIP: 856-73-46-215; NIP canonical repeat: 8567346215.",
    "Email: Synthetic.Person@Example.Test; Email canonical repeat: synthetic.person@example.test (same address).",
    "REGON: 192598474; REGON canonical repeat: 192598474.",
    "Telephone: +48 123 456 789; Telephone canonical repeat: +48-123-456-789.",
    "NRB: 11 1140 2004 0000 3002 1800 0002; NRB canonical repeat: 11114020040000300218000002.",
    "IBAN: PL61 1090 1014 0000 0712 1981 2874; IBAN canonical repeat: pl61109010140000071219812874.",
    "Near-miss PESEL: 44051401459.",
    "Near-miss NIP: 856-73-46-214.",
    "Near-miss email: synthetic.person@example.",
    "Near-miss telephone: +48 123 456 78.",
    "Near-miss NRB: 11 1140 2004 0000 3002 1800 0003.",
    "Near-miss IBAN: PL61 1090 1014 0000 0712 1981 2875.",
)
PII_DOCUMENT_TEXT = "\n".join(PII_DOCUMENT_LINES)
EXPECTED_ANONYMIZED_DOCUMENT_TEXT = "\n".join(
    (
        "PESEL: [PESEL_1]; PESEL canonical repeat: [PESEL_1].",
        "NIP: [NIP_1]; NIP canonical repeat: [NIP_1].",
        "Email: [EMAIL_1]; Email canonical repeat: [EMAIL_1] (same address).",
        "REGON: [REGON_1]; REGON canonical repeat: [REGON_1].",
        "Telephone: [TELEFON_1]; Telephone canonical repeat: [TELEFON_1].",
        "NRB: [NRB_1]; NRB canonical repeat: [NRB_1].",
        "IBAN: [IBAN_1]; IBAN canonical repeat: [IBAN_1].",
        "Near-miss PESEL: 44051401459.",
        "Near-miss NIP: 856-73-46-214.",
        "Near-miss email: synthetic.person@example.",
        "Near-miss telephone: +48 123 456 78.",
        "Near-miss NRB: 11 1140 2004 0000 3002 1800 0003.",
        "Near-miss IBAN: PL61 1090 1014 0000 0712 1981 2875.",
    )
)
KNOWN_MARKER_RESPONSE = (
    "Restored [PESEL_1], [NIP_1], [EMAIL_1], [REGON_1], [TELEFON_1], [NRB_1], [IBAN_1], and [UNKNOWN_1]."
)
EXPECTED_RESTORED_RESPONSE = (
    "Restored 44051401458, 856-73-46-215, Synthetic.Person@Example.Test, "
    "192598474, +48 123 456 789, 11 1140 2004 0000 3002 1800 0002, "
    "PL61 1090 1014 0000 0712 1981 2874, and [UNKNOWN_1]."
)


def make_pii_docx() -> bytes:
    document = Document()
    for line in PII_DOCUMENT_LINES:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pii_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    for index, line in enumerate(PII_DOCUMENT_LINES):
        page.insert_text((40, 60 + index * 28), line, fontsize=8)
    output = BytesIO()
    document.save(output)
    document.close()
    return output.getvalue()


def read_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def read_pdf_text(content: bytes) -> str:
    document = pymupdf.open(stream=content, filetype="pdf")
    try:
        return "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()
