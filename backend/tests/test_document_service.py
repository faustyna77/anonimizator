from io import BytesIO
import json
from types import SimpleNamespace
from uuid import UUID

from docx import Document
import pytest

from backend.app import document_service
from backend.app.auth import AccessContext
from backend.app.config import Settings
from backend.app.anonymization import DocumentAnonymizer, PhraseRule
from backend.app.document_repository import get_document_for_access_context
from backend.app.document_service import DocumentProcessingError, DocumentProcessingService
from backend.app.models import Office, Profile
from backend.app.office_rule_service import OfficeRuleInput, OfficeRuleService
from backend.app.storage import DocumentMappingCipher, DocumentStorage


class FakeS3Client:
    def __init__(self):
        self.objects = {}

    def put_object(self, **kwargs):
        self.objects[kwargs["Key"]] = kwargs["Body"]


class FakeRuleSnapshotService:
    def __init__(self, anonymizer):
        self.anonymizer = anonymizer
        self.calls = []

    def create_active_anonymizer(self, access_context):
        self.calls.append(access_context)
        return self.anonymizer


class FakeSession:
    def begin(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def close(self):
        pass


def _settings() -> Settings:
    return Settings(
        _env_file=None,
        s3_bucket="test-documents",
        s3_region="eu-central-1",
        aws_access_key_id="test-access-key",
        aws_secret_access_key="test-secret-key",
        document_mapping_encryption_key="MDEyMzQ1Njc4OWFiY2RlZjAxMjM0NTY3ODlhYmNkZWY=",
    )


def _context() -> AccessContext:
    return AccessContext(
        user_id="controlled-user",
        profile_id=UUID("00000000-0000-0000-0000-000000000001"),
        office_id=UUID("00000000-0000-0000-0000-000000000002"),
    )


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_service_ignores_a_disabled_rule_absent_from_the_active_snapshot(monkeypatch):
    fake_s3 = FakeS3Client()
    settings = _settings()
    created = []

    def create_processing(session, context, **kwargs):
        document = SimpleNamespace(id=kwargs["document_id"], status="processing", **kwargs)
        created.append(document)
        return document

    def mark_ready(session, context, document_id, **kwargs):
        document = created[0]
        document.status = "ready"
        document.encrypted_mapping = kwargs["encrypted_mapping"]
        return document

    monkeypatch.setattr(document_service.document_repository, "create_processing_document", create_processing)
    monkeypatch.setattr(document_service.document_repository, "mark_document_ready", mark_ready)
    snapshot_service = FakeRuleSnapshotService(DocumentAnonymizer())
    service = DocumentProcessingService(
        session_factory=FakeSession,
        storage=DocumentStorage(settings, client=fake_s3),
        mapping_cipher=DocumentMappingCipher.from_settings(settings),
        office_rule_service=snapshot_service,
    )

    result = service.process_upload(
        _context(),
        content=_docx_bytes("Jan Kowalski"),
        original_filename="synthetic.docx",
        document_format="docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result_key = f"offices/{_context().office_id}/documents/{result.document_id}/anonymized"
    assert Document(BytesIO(fake_s3.objects[result_key])).paragraphs[0].text == "Jan Kowalski"
    assert json.loads(DocumentMappingCipher.from_settings(settings).decrypt(created[0].encrypted_mapping)) == {}


def test_service_loads_one_active_rule_snapshot_before_transforming_a_document(monkeypatch):
    fake_s3 = FakeS3Client()
    settings = _settings()
    created = []

    def create_processing(session, context, **kwargs):
        document = SimpleNamespace(id=kwargs["document_id"], status="processing", **kwargs)
        created.append(document)
        return document

    def mark_ready(session, context, document_id, **kwargs):
        document = created[0]
        document.status = "ready"
        document.encrypted_mapping = kwargs["encrypted_mapping"]
        return document

    monkeypatch.setattr(document_service.document_repository, "create_processing_document", create_processing)
    monkeypatch.setattr(document_service.document_repository, "mark_document_ready", mark_ready)
    snapshot_service = FakeRuleSnapshotService(
        DocumentAnonymizer(phrase_rules=(PhraseRule(rule_id="office-rule", phrase="Jan Kowalski", marker_label="KLIENT"),))
    )
    service = DocumentProcessingService(
        session_factory=FakeSession,
        storage=DocumentStorage(settings, client=fake_s3),
        mapping_cipher=DocumentMappingCipher.from_settings(settings),
        office_rule_service=snapshot_service,
    )

    result = service.process_upload(
        _context(),
        content=_docx_bytes("Jan Kowalski / Jan Kowalski"),
        original_filename="synthetic.docx",
        document_format="docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert snapshot_service.calls == [_context()]
    result_key = f"offices/{_context().office_id}/documents/{result.document_id}/anonymized"
    assert Document(BytesIO(fake_s3.objects[result_key])).paragraphs[0].text == "[KLIENT_1] / [KLIENT_1]"
    assert json.loads(DocumentMappingCipher.from_settings(settings).decrypt(created[0].encrypted_mapping)) == {
        "[KLIENT_1]": "Jan Kowalski"
    }


@pytest.mark.postgresql
def test_service_excludes_foreign_persisted_rules_from_document_artifact_and_mapping(
    migrated_session_factory,
):
    settings = _settings()
    fake_s3 = FakeS3Client()
    with migrated_session_factory.begin() as session:
        first_office = Office(name="First processing office")
        second_office = Office(name="Second processing office")
        first_profile = Profile(user_id="first-processing-user", office=first_office)
        second_profile = Profile(user_id="second-processing-user", office=second_office)
        session.add_all([first_profile, second_profile])
        session.flush()
        first_context = AccessContext("first-processing-user", first_profile.id, first_office.id)
        second_context = AccessContext("second-processing-user", second_profile.id, second_office.id)

    rule_service = OfficeRuleService(session_factory=migrated_session_factory)
    rule_service.create_rule(
        first_context,
        OfficeRuleInput(
            kind="phrase",
            pattern="FIRST-OFFICE-SECRET",
            marker_label="FIRST_SECRET",
            enabled=True,
        ),
    )
    rule_service.create_rule(
        second_context,
        OfficeRuleInput(
            kind="phrase",
            pattern="SECOND-OFFICE-SECRET",
            marker_label="SECOND_SECRET",
            enabled=True,
        ),
    )
    cipher = DocumentMappingCipher.from_settings(settings)
    service = DocumentProcessingService(
        session_factory=migrated_session_factory,
        storage=DocumentStorage(settings, client=fake_s3),
        mapping_cipher=cipher,
        office_rule_service=rule_service,
    )

    result = service.process_upload(
        second_context,
        content=_docx_bytes("FIRST-OFFICE-SECRET / SECOND-OFFICE-SECRET"),
        original_filename="second-office.docx",
        document_format="docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result_key = f"offices/{second_context.office_id}/documents/{result.document_id}/anonymized"
    assert Document(BytesIO(fake_s3.objects[result_key])).paragraphs[0].text == (
        "FIRST-OFFICE-SECRET / [SECOND_SECRET_1]"
    )
    with migrated_session_factory() as session:
        persisted_document = get_document_for_access_context(session, second_context, result.document_id)

    assert persisted_document is not None
    assert persisted_document.encrypted_mapping is not None
    assert json.loads(cipher.decrypt(persisted_document.encrypted_mapping)) == {
        "[SECOND_SECRET_1]": "SECOND-OFFICE-SECRET"
    }


def test_service_persists_synthetic_original_result_and_encrypted_mapping_with_fake_s3(monkeypatch):
    fake_s3 = FakeS3Client()
    settings = _settings()
    created = []

    def create_processing(session, context, **kwargs):
        document = SimpleNamespace(id=kwargs["document_id"], status="processing", **kwargs)
        created.append(document)
        return document

    def mark_ready(session, context, document_id, **kwargs):
        document = created[0]
        document.status = "ready"
        document.anonymized_object_key = kwargs["anonymized_object_key"].value
        document.encrypted_mapping = kwargs["encrypted_mapping"]
        return document

    monkeypatch.setattr(document_service.document_repository, "create_processing_document", create_processing)
    monkeypatch.setattr(document_service.document_repository, "mark_document_ready", mark_ready)
    service = DocumentProcessingService(
        session_factory=FakeSession,
        storage=DocumentStorage(settings, client=fake_s3),
        mapping_cipher=DocumentMappingCipher.from_settings(settings),
        office_rule_service=FakeRuleSnapshotService(DocumentAnonymizer()),
    )

    result = service.process_upload(
        _context(),
        content=_docx_bytes("Synthetic identifier 44051401458."),
        original_filename="synthetic.docx",
        document_format="docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.document_format == "docx"
    assert result.original_filename == "synthetic.docx"
    assert result.status == "ready"
    assert created[0].status == "ready"
    assert len(fake_s3.objects) == 2
    assert f"offices/{_context().office_id}/documents/{result.document_id}/original" in fake_s3.objects
    result_key = f"offices/{_context().office_id}/documents/{result.document_id}/anonymized"
    assert result_key in fake_s3.objects
    decrypted = DocumentMappingCipher.from_settings(settings).decrypt(created[0].encrypted_mapping)
    assert json.loads(decrypted) == {"[PESEL_1]": "44051401458"}
    assert b"44051401458" not in fake_s3.objects[result_key]


def test_service_marks_document_failed_after_a_controlled_format_error(monkeypatch):
    fake_s3 = FakeS3Client()
    settings = _settings()
    created = []

    def create_processing(session, context, **kwargs):
        document = SimpleNamespace(id=kwargs["document_id"], status="processing", **kwargs)
        created.append(document)
        return document

    def mark_failed(session, context, document_id):
        created[0].status = "failed"
        return created[0]

    monkeypatch.setattr(document_service.document_repository, "create_processing_document", create_processing)
    monkeypatch.setattr(document_service.document_repository, "mark_document_failed", mark_failed)
    service = DocumentProcessingService(
        session_factory=FakeSession,
        storage=DocumentStorage(settings, client=fake_s3),
        mapping_cipher=DocumentMappingCipher.from_settings(settings),
        office_rule_service=FakeRuleSnapshotService(DocumentAnonymizer()),
    )

    try:
        service.process_upload(
            _context(),
            content=b"not-a-pdf",
            original_filename="synthetic.pdf",
            document_format="pdf",
            content_type="application/pdf",
        )
    except DocumentProcessingError as error:
        assert str(error) == "The PDF file could not be processed."
    else:
        raise AssertionError("invalid document unexpectedly processed")

    assert created[0].status == "failed"
    assert len(fake_s3.objects) == 1
