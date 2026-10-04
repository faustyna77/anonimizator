"""Format-specific document transformation adapters with no OCR fallback."""
from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from typing import Any

import pymupdf
from docx import Document
from docx.table import Table, _Cell

from backend.app.anonymization import DocumentAnonymizer


class DocumentProcessingError(Exception):
    """A client-safe processing failure that never includes source document content."""


class UnsupportedDocumentError(DocumentProcessingError):
    """The uploaded bytes cannot be processed as the declared document format."""


def anonymize_docx(content: bytes) -> tuple[bytes, dict[str, str]]:
    """Replace identifiers in DOCX body, table, header, and footer paragraphs."""
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise UnsupportedDocumentError("The DOCX file could not be processed.") from error

    anonymizer = DocumentAnonymizer()
    for paragraph in _document_paragraphs(document):
        _anonymize_docx_paragraph(paragraph, anonymizer)

    result = BytesIO()
    try:
        document.save(result)
    except Exception as error:
        raise DocumentProcessingError("The DOCX file could not be processed.") from error
    return result.getvalue(), anonymizer.anonymize("").mapping


def _anonymize_docx_paragraph(paragraph: Any, anonymizer: DocumentAnonymizer) -> None:
    """Replace matched text in-place so existing runs retain their character formatting."""
    runs = list(paragraph.runs)
    source = "".join(run.text for run in runs)
    if not source:
        return

    anonymized = anonymizer.anonymize(source)
    if anonymized.text == source:
        return

    replacements = anonymized.replacements
    if not replacements:
        raise DocumentProcessingError("The DOCX file could not be processed.")

    run_ranges = _docx_run_ranges(runs)
    for start, end, marker in reversed(replacements):
        start_run, start_offset = _docx_run_at_position(run_ranges, start, is_end=False)
        end_run, end_offset = _docx_run_at_position(run_ranges, end, is_end=True)
        if start_run == end_run:
            run = runs[start_run]
            run.text = run.text[:start_offset] + marker + run.text[end_offset:]
            continue

        runs[start_run].text = runs[start_run].text[:start_offset] + marker
        for index in range(start_run + 1, end_run):
            runs[index].text = ""
        runs[end_run].text = runs[end_run].text[end_offset:]


def _docx_run_ranges(runs: list[Any]) -> list[tuple[int, int]]:
    position = 0
    ranges = []
    for run in runs:
        next_position = position + len(run.text)
        ranges.append((position, next_position))
        position = next_position
    return ranges


def _docx_run_at_position(
    run_ranges: list[tuple[int, int]], position: int, *, is_end: bool
) -> tuple[int, int]:
    for index, (start, end) in enumerate(run_ranges):
        if start <= position < end or (is_end and position == end):
            return index, position - start
    raise DocumentProcessingError("The DOCX file could not be processed.")


def _document_paragraphs(document: Document) -> Iterable[Any]:
    for paragraph in document.paragraphs:
        yield paragraph
    yield from _table_paragraphs(document.tables)
    for section in document.sections:
        for header_or_footer in (section.header, section.footer):
            for paragraph in header_or_footer.paragraphs:
                yield paragraph
            yield from _table_paragraphs(header_or_footer.tables)


def _table_paragraphs(tables: Iterable[Table]) -> Iterable[Any]:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _cell_paragraphs(cell)


def _cell_paragraphs(cell: _Cell) -> Iterable[Any]:
    for paragraph in cell.paragraphs:
        yield paragraph
    yield from _table_paragraphs(cell.tables)


def anonymize_pdf(content: bytes) -> tuple[bytes, dict[str, str]]:
    """Redact identifiers in-place while retaining the source PDF's pages and layout."""
    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception as error:
        raise UnsupportedDocumentError("The PDF file could not be processed.") from error

    try:
        if document.page_count == 0 or not any(page.get_text("text").strip() for page in document):
            raise DocumentProcessingError("PDF text layer is unavailable; OCR is not supported in this MVP.")

        anonymizer = DocumentAnonymizer()
        for page in document:
            _redact_pdf_page(page, anonymizer)

        result = BytesIO()
        document.save(result, garbage=4, deflate=True)
        return result.getvalue(), anonymizer.anonymize("").mapping
    except DocumentProcessingError:
        raise
    except Exception as error:
        raise DocumentProcessingError("The PDF file could not be processed.") from error
    finally:
        document.close()


def _redact_pdf_page(page: pymupdf.Page, anonymizer: DocumentAnonymizer) -> None:
    source = page.get_text("text")
    if not source.strip():
        return

    replacements = anonymizer.anonymize(source).replacements
    annotations_added = False
    seen_sources: set[tuple[str, str]] = set()
    for start, end, marker in replacements:
        identifier = source[start:end]
        source_key = (identifier, marker)
        if source_key in seen_sources:
            continue
        seen_sources.add(source_key)

        rectangles = page.search_for(identifier)
        if not rectangles:
            raise DocumentProcessingError("The PDF file could not be processed.")
        for rectangle in rectangles:
            page.add_redact_annot(
                rectangle,
                text=marker,
                fontname="helv",
                fontsize=_redaction_font_size(rectangle, marker),
                fill=(1, 1, 1),
                text_color=(0, 0, 0),
                cross_out=False,
            )
            annotations_added = True

    if annotations_added:
        page.apply_redactions()


def _redaction_font_size(rectangle: pymupdf.Rect, marker: str) -> float:
    """Fit the marker into the original identifier's redaction rectangle."""
    size = max(4.0, min(rectangle.height * 0.8, 12.0))
    marker_width = pymupdf.get_text_length(marker, fontname="helv", fontsize=size)
    if marker_width > rectangle.width:
        size = max(4.0, size * rectangle.width / marker_width)
    return size
